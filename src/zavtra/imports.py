"""Ограниченный CSV/XLSX импорт. Валидация завершена до изменения БД."""

import csv
import io
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from zavtra.domain import Sale, UserError, number

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 50000
ALIASES = {"дата": "date", "блюдо": "dish", "количество": "quantity"}


def table(content: bytes, filename: str) -> list[dict[str, object]]:
    if len(content) > MAX_BYTES:
        raise UserError("Файл больше 5 МБ.")
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".csv":
            text = content.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            rows = list(csv.reader(io.StringIO(text), dialect))
        elif ext == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise UserError("Распакованный XLSX больше 50 МБ.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                rows = []
                for row in workbook.active.iter_rows(values_only=True):
                    rows.append(list(row))
                    if len(rows) > MAX_ROWS + 1:
                        raise UserError(f"Максимум {MAX_ROWS} строк.")
            finally:
                workbook.close()
        else:
            raise UserError("Отправьте CSV в UTF-8 или XLSX.")
    except (UnicodeDecodeError, csv.Error, zipfile.BadZipFile, OSError, ValueError) as exc:
        if isinstance(exc, UserError):
            raise
        raise UserError("Не удалось прочитать файл. Проверьте формат CSV/XLSX.") from None
    if len(rows) < 2:
        raise UserError("В файле нужны заголовок и хотя бы одна строка данных.")
    if len(rows) > MAX_ROWS + 1:
        raise UserError(f"Максимум {MAX_ROWS} строк.")
    headers = [str(v or "").strip().lower() for v in rows[0]]
    headers = [ALIASES.get(v, v) for v in headers]
    if not all(headers) or len(set(headers)) != len(headers):
        raise UserError("Заголовки колонок должны быть непустыми и уникальными.")
    result = []
    for index, row in enumerate(rows[1:], 2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        if len(row) != len(headers):
            raise UserError(f"Строка {index}: число ячеек отличается от заголовка.")
        result.append(dict(zip(headers, row, strict=True)) | {"_line": index})
    if not result:
        raise UserError("Нет строк данных.")
    return result


def required(rows: list[dict], columns: set[str]) -> None:
    missing = columns - rows[0].keys()
    if missing:
        raise UserError("Не хватает колонок: " + ", ".join(sorted(missing)))


def parse_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise UserError("дата должна быть ГГГГ-ММ-ДД или ДД.ММ.ГГГГ")


def sales_file(content: bytes, filename: str, today: date) -> list[Sale]:
    rows = table(content, filename)
    required(rows, {"date", "dish", "quantity"})
    result, errors = [], []
    for row in rows:
        try:
            day = parse_date(row["date"])
            if day > today:
                raise UserError("дата в будущем")
            if day < today - timedelta(days=730):
                raise UserError("дата старше двух лет")
            dish = str(row["dish"] or "").strip()
            if not dish or len(dish) > 100:
                raise UserError("название блюда: от 1 до 100 символов")
            quantity = number(row["quantity"], "Количество")
            if quantity != quantity.to_integral_value():
                raise UserError("количество порций должно быть целым")
            result.append(Sale(day, dish, int(quantity)))
        except UserError as exc:
            errors.append(f"Строка {row['_line']}: {exc}")
    if errors:
        suffix = f"\nВсего ошибок: {len(errors)}. Файл не загружен."
        raise UserError("\n".join(errors[:15]) + suffix)
    if len({row.dish for row in result}) > 100 or len({row.day for row in result}) > 365:
        raise UserError("Максимум 100 блюд и 365 дат в файле.")
    return result
