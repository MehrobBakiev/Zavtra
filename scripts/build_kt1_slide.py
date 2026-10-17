"""Слайд по шаблону КТ-1. Зависимость: python-pptx==1.0.2.

Запуск: python scripts/build_kt1_slide.py
PDF: libreoffice --headless --convert-to pdf --outdir Documents_Zavtra \
    Documents_Zavtra/KT1_Zavtra.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
GRAY = RGBColor.from_string("595959")
BLACK = RGBColor.from_string("000000")
LINK = RGBColor.from_string("009FB5")


def textbox(slide, x, y, width, height, text, size=20, color=GRAY, url=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.line_spacing = 1.12
        paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = line
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.color.rgb = color
        if url:
            run.font.underline = True
            run.hyperlink.address = url
    return box


def main():
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = "Завтра — контрольная точка 1"
    presentation.core_properties.author = "Бакиев Мехроб Мехрзодович"
    presentation.core_properties.subject = "Проектный семинар · 17.10.2026"
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    textbox(slide, 0.6, 0.65, 4.7, 0.6, "Завтра", 30, BLACK)
    textbox(slide, 0.6, 1.65, 4.6, 0.85, "Автор: Бакиев Мехроб\nМехрзодович", 20)
    textbox(slide, 0.6, 2.65, 4.7, 0.45, "Описание проекта:", 22)

    base = "https://github.com/MehrobBakiev/Zavtra"
    documents = [
        ("Одностраничник", base + "/blob/main/Documents_Zavtra/one_pager_zavtra_2.pdf"),
        ("PRD", base + "/blob/main/Documents_Zavtra/2.1_PRD_Zavtra.pdf"),
        ("SRS", base + "/blob/main/Documents_Zavtra/2.2_SRS_Zavtra.pdf"),
        ("User Stories", base + "/blob/main/Documents_Zavtra/2.3_UserStories_Zavtra.pdf"),
        ("RTM", base + "/blob/main/Documents_Zavtra/2.4_RTM_Zavtra.pdf"),
        ("Git-репозиторий", base),
    ]
    for index, (label, url) in enumerate(documents):
        y = 3.3 + index * 0.4
        textbox(slide, 0.75, y, 0.2, 0.38, "•", 20)
        textbox(slide, 1.1, y, 4.0, 0.38, label, 20, LINK, url)

    textbox(slide, 5.55, 0.67, 7.0, 0.5, "Что было сделано:", 23)
    bullets = [
        "Реализован сценарий Telegram-бота:\nпродажи → остатки → рекомендация →\nподтверждение заказа.",
        "Добавлены CSV/XLSX, техкарты,\nручная правка заказа и хранение данных\nв SQLite / PostgreSQL.",
        "Автоматические тесты: 31 из 31.\nПроверки GitHub Actions успешны\nна Python 3.12 и 3.14.",
        "Код и документы опубликованы в GitHub.\nПодготовлены демоданные и инструкция\nзапуска.",
    ]
    for index, text in enumerate(bullets):
        y = 1.45 + index * 1.17
        textbox(slide, 5.65, y, 0.3, 0.4, "•", 22)
        textbox(slide, 6.05, y, 6.65, 1.07, text, 20)
    textbox(slide, 6.05, 6.45, 6.65, 0.35, "Пилот на реальной точке — следующий этап.", 14)
    textbox(slide, 0.75, 7.03, 3.0, 0.3, "17.10.2026", 14)

    output = ROOT / "Documents_Zavtra/KT1_Zavtra.pptx"
    presentation.save(output)
    print(output)


if __name__ == "__main__":
    main()
